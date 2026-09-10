from __future__ import annotations

import html
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .config import NegotiationSettings
from .schemas import BriefArtifactDTO, BriefContentDTO, ContactDTO, MessageDTO, PromptTemplateDTO


class BriefService:
    def __init__(self, settings: NegotiationSettings):
        self.settings = settings

    def build_content(self, contact: ContactDTO, messages: list[MessageDTO], prompt: PromptTemplateDTO) -> BriefContentDTO:
        texts = [message.text for message in messages if message.text.strip()]
        sample = " ".join(texts[-12:])
        questions = [
            "Какая задача сейчас самая срочная?",
            "Как вы сейчас решаете эту задачу без GramLead?",
            "Что должно измениться после внедрения, чтобы встреча считалась успешной?",
        ]
        buying_signals = [text[:180] for text in texts if any(word in text.lower() for word in ("нужно", "хочу", "интерес", "куп", "заказ", "цена", "стоимость"))][:5]
        if not buying_signals and texts:
            buying_signals = [texts[-1][:180]]
        return BriefContentDTO(
            client_summary=f"{contact.title}. Источник: {contact.source_title or contact.source_id or 'GramLead'}. Сообщений в базе: {contact.messages_count}.",
            conversation_summary=(sample[:1200] or "Для контакта мало текстового контекста; перед встречей стоит уточнить задачу и текущий процесс."),
            pain_hypotheses=[
                "Клиент может терять лиды и важные сигналы в длинных переписках.",
                "Подготовка к переговорам занимает ручное время.",
                "Нужна понятная выжимка по контактам и источникам перед продажей.",
            ],
            buying_signals=buying_signals,
            risk_flags=[
                "Мало контекста по бюджету и срокам.",
                "Нужно не обещать автоматический результат без проверки источников данных.",
            ],
            objections=[
                "У нас уже есть чаты и таблицы.",
                "Неясно, насколько быстро окупится внедрение.",
            ],
            arguments=[
                "GramLead превращает историю сообщений в структурированные сигналы продаж.",
                "Менеджер получает бриф до разговора, а не перечитывает переписку вручную.",
                "Можно быстро выгружать материалы по каждому контакту в понятные форматы.",
            ],
            questions=questions,
            next_step="Провести короткую диагностическую встречу и выбрать 2-3 источника для пилотного анализа.",
            follow_up_message=f"Здравствуйте! Перед встречей посмотрел контекст по вашему запросу. Предлагаю обсудить, где сейчас теряются лиды и какие источники важнее всего разобрать первыми. Подойдет короткий созвон на 20 минут?",
            manager_checklist=[
                "Открыть последние сообщения контакта.",
                "Проверить источник и дату последнего сообщения.",
                "Выбрать один главный вопрос для начала разговора.",
                "Не обещать интеграции, которых нет в текущем контуре.",
            ],
            data_confidence="high" if len(messages) >= 20 else "medium" if messages else "low",
            missing_context=["бюджет", "срок принятия решения", "кто принимает решение"],
            recommended_offer="Пилот GramLead на ограниченном наборе источников с брифами по контактам.",
            do_not_say=["Гарантируем продажи без работы менеджера.", "Все интеграции уже готовы."],
        )

    def write_artifacts(self, brief_id: str, contact: ContactDTO, messages: list[MessageDTO], content: BriefContentDTO) -> list[BriefArtifactDTO]:
        root = self.settings.artifacts_path / brief_id
        root.mkdir(parents=True, exist_ok=True)
        payload = {
            "brief_id": brief_id,
            "contact": contact.model_dump(),
            "messages": [message.model_dump() for message in messages],
            "content": content.model_dump(),
        }
        artifacts = []
        json_path = root / "brief.json"
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        artifacts.append(self._artifact("json", json_path))
        md_path = root / "brief.md"
        md_path.write_text(self._to_markdown(contact, content), encoding="utf-8")
        artifacts.append(self._artifact("markdown", md_path))
        html_path = root / "brief.html"
        html_path.write_text(self._to_html(contact, content), encoding="utf-8")
        artifacts.append(self._artifact("html", html_path))
        docx_path = root / "brief.docx"
        self._write_docx(docx_path, contact, content)
        artifacts.append(self._artifact("docx", docx_path))
        input_path = root / "brief-input-context.json"
        input_path.write_text(json.dumps({"contact": contact.model_dump(), "messages": [m.model_dump() for m in messages]}, ensure_ascii=False, indent=2), encoding="utf-8")
        artifacts.append(self._artifact("input_context", input_path))
        return artifacts

    def _artifact(self, kind: str, path: Path) -> BriefArtifactDTO:
        return BriefArtifactDTO(kind=kind, path=str(path), url=f"{self.settings.public_base_url}/files/{path.parent.name}/{path.name}")

    @staticmethod
    def _to_markdown(contact: ContactDTO, content: BriefContentDTO) -> str:
        lines = [f"# Бриф к переговорам: {contact.title}", "", "## Резюме", content.client_summary, "", "## История", content.conversation_summary]
        sections = [
            ("Гипотезы болей", content.pain_hypotheses),
            ("Сигналы покупки", content.buying_signals),
            ("Риски", content.risk_flags),
            ("Возражения", content.objections),
            ("Аргументы", content.arguments),
            ("Вопросы", content.questions),
            ("Чеклист менеджера", content.manager_checklist),
            ("Что не говорить", content.do_not_say),
        ]
        for title, items in sections:
            lines.extend(["", f"## {title}"])
            lines.extend([f"- {item}" for item in items] or ["- нет данных"])
        lines.extend(["", "## Следующий шаг", content.next_step, "", "## Сообщение после встречи", content.follow_up_message])
        return "\n".join(lines)

    def _to_html(self, contact: ContactDTO, content: BriefContentDTO) -> str:
        body = html.escape(self._to_markdown(contact, content)).replace("\n", "<br>")
        return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(contact.title)}</title></head><body>{body}</body></html>"

    def _write_docx(self, path: Path, contact: ContactDTO, content: BriefContentDTO) -> None:
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10.5)
        doc.add_heading(f"Бриф к переговорам: {contact.title}", 0)
        for title, value in [
            ("Резюме", content.client_summary),
            ("История", content.conversation_summary),
            ("Следующий шаг", content.next_step),
            ("Сообщение после встречи", content.follow_up_message),
        ]:
            doc.add_heading(title, level=1)
            doc.add_paragraph(value)
        for title, items in [
            ("Гипотезы болей", content.pain_hypotheses),
            ("Сигналы покупки", content.buying_signals),
            ("Риски", content.risk_flags),
            ("Вопросы", content.questions),
            ("Чеклист", content.manager_checklist),
        ]:
            doc.add_heading(title, level=1)
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
        doc.save(path)

