from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(
    "/Users/aiapi/Desktop/gramlead-backfront/gramlead-main-new-optimize/"
    "docs/justai-negotiation-assistant-gramlead/word-cases"
)

BLUE = "1F4D78"
LIGHT = "F2F4F7"
INK = "111827"
MUTED = "64748B"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True)
        shade(table.rows[0].cells[index], LIGHT)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string(BLUE)


def build_doc(filename: str, title: str, subtitle: str, sections: list[tuple[str, list[tuple]]]) -> Path:
    doc = Document()
    configure_styles(doc)

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor.from_string(INK)
    title_paragraph.paragraph_format.space_after = Pt(2)

    subtitle_paragraph = doc.add_paragraph()
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle_paragraph.paragraph_format.space_after = Pt(10)

    for section_title, body in sections:
        doc.add_heading(section_title, level=1)
        for block in body:
            kind = block[0]
            if kind == "p":
                paragraph = doc.add_paragraph(block[1])
                paragraph.paragraph_format.space_after = Pt(5)
            elif kind == "bullets":
                add_bullets(doc, block[1])
            elif kind == "table":
                add_table(doc, block[1], block[2], block[3])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "GramLead · отдельный Docker-модуль подготовки к переговорам · источник идеи: Just AI marketplace"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename
    doc.save(path)
    return path


COMMON_SOURCE = (
    "Just AI описывает ассистента как агента, который собирает контекст по клиенту "
    "и формирует краткую выжимку: боли, вопросы, аргументы, риски и следующий шаг. "
    "В GramLead этот сценарий строится поверх DuckDB, JSONL, HTML/XLSX/DOCX artifacts "
    "и LLM-аналитики без повторного чтения Telegram."
)


def main() -> None:
    docs = [
        build_doc(
            "01-case-sales-manager-negotiation-brief.docx",
            "Кейс 1. Менеджер готовится к переговорам с клиентом",
            "Как отдельный Docker-модуль GramLead превращает историю сообщений и локальные artifacts в бриф перед встречей.",
            [
                (
                    "Сценарий",
                    [
                        ("p", COMMON_SOURCE),
                        (
                            "p",
                            "Менеджер выбирает контакт или компанию, нажимает “Подготовить переговоры” "
                            "и получает готовый бриф: кто клиент, что уже писал, какие боли проявлял, "
                            "какие аргументы использовать и каким вопросом открыть встречу.",
                        ),
                    ],
                ),
                (
                    "Входные данные",
                    [
                        (
                            "table",
                            ["Источник", "Что используется"],
                            [
                                ["GramLead DuckDB", "контакты, сообщения, источники, даты, авторы"],
                                ["Файлы контакта", "JSONL/HTML/XLSX история сообщений"],
                                ["OpenRouter", "LLM-анализ намерений, болей и вероятности покупки"],
                                ["GramLead artifacts", "локальные ссылки на историю, анализ и документы контакта"],
                            ],
                            [2.0, 4.2],
                        )
                    ],
                ),
                (
                    "Результат",
                    [
                        (
                            "bullets",
                            [
                                "Краткое резюме клиента на 5-10 строк.",
                                "Гипотезы потребностей и болей.",
                                "3-5 аргументов под продукт GramLead.",
                                "Риски и возможные возражения.",
                                "Сценарий первого вопроса и следующий шаг после встречи.",
                            ],
                        )
                    ],
                ),
            ],
        ),
        build_doc(
            "02-case-rop-systemic-sales-preparation.docx",
            "Кейс 2. РОП делает подготовку к встречам системной",
            "Как руководитель продаж контролирует качество подготовки менеджеров по всем клиентам.",
            [
                (
                    "Задача РОПа",
                    [
                        (
                            "p",
                            "У РОПа нет прозрачности: менеджеры готовятся по-разному, часть контекста остается "
                            "в Telegram, часть в локальных файлах, часть в голове. Модуль подготовки к переговорам "
                            "стандартизирует бриф и делает подготовку измеримой.",
                        )
                    ],
                ),
                (
                    "Контрольные метрики",
                    [
                        (
                            "table",
                            ["Метрика", "Зачем нужна"],
                            [
                                ["Контакты с брифом", "понимать покрытие базы перед звонками"],
                                ["Среднее число сообщений в анализе", "оценить глубину подготовки"],
                                ["Дата последнего контакта", "не идти на встречу со старым контекстом"],
                                ["Риск/теплота клиента", "приоритизировать встречи"],
                                ["Следующий шаг", "контролировать управляемость сделки"],
                            ],
                            [2.4, 3.8],
                        )
                    ],
                ),
                (
                    "Как это ложится на GramLead",
                    [
                        (
                            "bullets",
                            [
                                "Отдельный контейнер читает базу GramLead только на чтение.",
                                "Результаты пишет в собственную таблицу/папку artifacts.",
                                "Основной GramLead не блокируется тяжелыми LLM-задачами.",
                                "РОП видит список брифов, статусы готовности и ссылки на локальные артефакты.",
                            ],
                        )
                    ],
                ),
            ],
        ),
        build_doc(
            "03-case-telegram-history-to-meeting-brief.docx",
            "Кейс 3. Из истории Telegram в переговорный бриф",
            "Как использовать уже накопленные сообщения GramLead для подготовки без ручного перечитывания переписки.",
            [
                (
                    "Проблема",
                    [
                        (
                            "p",
                            "В Telegram-источниках много сигналов: вопросы, жалобы, интерес к продукту, "
                            "бюджет, срочность. Ручное чтение занимает часы и часто приводит к потере важных деталей.",
                        )
                    ],
                ),
                (
                    "Пайплайн",
                    [
                        (
                            "table",
                            ["Шаг", "Что делает модуль"],
                            [
                                ["1. Выбор контакта", "берет контакт из DuckDB и локального индекса GramLead"],
                                ["2. Сбор сообщений", "находит релевантную историю и последние реплики"],
                                ["3. Нормализация", "убирает шум, системные сообщения, дубли"],
                                ["4. LLM-анализ", "выделяет боли, мотивы, возражения, потенциальную ценность"],
                                ["5. Бриф", "создает Markdown/DOCX/HTML результат для менеджера"],
                            ],
                            [1.6, 4.6],
                        )
                    ],
                ),
                (
                    "Формат брифа",
                    [
                        (
                            "bullets",
                            [
                                "Кто клиент и откуда пришел.",
                                "Краткая история коммуникации.",
                                "Что может купить и почему.",
                                "Что нельзя говорить или обещать.",
                                "Вопросы для выявления потребности.",
                                "Письмо/сообщение после встречи.",
                            ],
                        )
                    ],
                ),
            ],
        ),
        build_doc(
            "04-case-separate-docker-module-architecture.docx",
            "Кейс 4. Отдельный Docker-контейнер на базе GramLead",
            "Архитектурный вариант реализации ассистента подготовки к переговорам как самостоятельного модуля.",
            [
                (
                    "Почему отдельный контейнер",
                    [
                        (
                            "p",
                            "Ассистент подготовки к переговорам должен быть изолирован от Telegram-сканирования "
                            "и основного UI. Он читает готовую базу GramLead, создает брифы и не ломает поток "
                            "импорта/синхронизации.",
                        )
                    ],
                ),
                (
                    "Компоненты",
                    [
                        (
                            "table",
                            ["Компонент", "Ответственность"],
                            [
                                ["negotiation-front", "UI выбора контакта, генерации и просмотра брифов"],
                                ["negotiation-back", "API, задания, доступ к DuckDB snapshot"],
                                ["negotiation-worker", "LLM-анализ и генерация DOCX/HTML"],
                                ["artifacts storage", "брифы, логи, файлы экспорта"],
                                ["shared read DB", "read-only доступ к GramLead DuckDB/snapshot"],
                            ],
                            [2.0, 4.2],
                        )
                    ],
                ),
                (
                    "Безопасность",
                    [
                        (
                            "bullets",
                            [
                                "DuckDB открывается на чтение или через snapshot, чтобы не блокировать основной процесс.",
                                "Секреты OpenRouter хранятся в настройках модуля, не в DOCX/HTML результатах.",
                                "В брифе показываются только нужные менеджеру данные.",
                                "Каждый LLM-запрос логируется с ссылкой на исходный контакт.",
                            ],
                        )
                    ],
                ),
            ],
        ),
    ]
    for doc_path in docs:
        print(doc_path)


if __name__ == "__main__":
    main()
